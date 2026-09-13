from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "9월8일_G조_발표_스크립트.md"
OUTPUT = ROOT / "docs" / "9월8일_G조_발표_스크립트.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Noto Sans CJK KR"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.3

for style_name, size, color in (
    ("Title", 25, "2468CC"),
    ("Heading 1", 20, "2468CC"),
    ("Heading 2", 17, "2468CC"),
):
    style = styles[style_name]
    style.font.name = "Noto Sans CJK KR"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)

add_page_number(section.footer.paragraphs[0])

lines = SOURCE.read_text(encoding="utf-8").splitlines()
slide_started = False

for raw in lines:
    line = raw.strip()
    if not line or line == "---":
        continue
    if line.startswith("# "):
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line[2:])
        continue
    if line.startswith("예상 발표 시간:") or line.startswith("발표자:"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line.replace("  ", ""))
        continue
    if line.startswith("## "):
        if slide_started:
            doc.add_page_break()
        slide_started = True
        table = doc.add_table(rows=1, cols=1)
        table.autofit = True
        cell = table.cell(0, 0)
        set_cell_shading(cell, "EAF2FF")
        p = cell.paragraphs[0]
        p.style = styles["Heading 1"]
        p.add_run(line[3:])
        continue
    if line.startswith("> 다음 페이지:"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        run = p.add_run("▶ " + line.split(":", 1)[1].strip().strip('“”\"'))
        run.bold = True
        run.font.color.rgb = RGBColor(93, 74, 199)
        continue
    p = doc.add_paragraph()
    p.add_run(line)

doc.core_properties.title = "9월 8일 G조 발표 스크립트"
doc.core_properties.subject = "M.A.R.S 프로젝트 페이지별 발표 대본"
doc.core_properties.author = "G조"
doc.save(OUTPUT)
print(OUTPUT)
