"""Safe text extraction for files explicitly attached to cloud nodes."""

from pathlib import Path

from sqlalchemy.orm import Session

from .. import models
from ..config import UPLOAD_DIR

MAX_FILE_CHARS = 30_000
MAX_FILES = 10


class CloudFileError(ValueError):
    pass


def _joined(parts: list[str]) -> str:
    return "\n".join(parts)[:MAX_FILE_CHARS]


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            return _joined([(page.extract_text() or "") for page in PdfReader(str(path)).pages])
        if suffix == ".docx":
            from docx import Document

            document = Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            for table in document.tables:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return _joined(parts)
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(str(path), read_only=True, data_only=True)
            parts: list[str] = []
            for worksheet in workbook.worksheets[:10]:
                parts.append(f"[Sheet: {worksheet.title}]")
                for row in worksheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        parts.append(" | ".join(values))
                    if sum(map(len, parts)) >= MAX_FILE_CHARS:
                        break
            return _joined(parts)
        if suffix == ".pptx":
            from pptx import Presentation

            parts = []
            for index, slide in enumerate(Presentation(str(path)).slides, start=1):
                parts.append(f"[Slide {index}]")
                parts.extend(
                    shape.text for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False) and shape.text
                )
            return _joined(parts)
        return path.read_bytes()[:MAX_FILE_CHARS].decode("utf-8", errors="replace")
    except Exception as exc:
        raise CloudFileError(f"Cannot extract text from {path.name}: {type(exc).__name__}") from exc


def cloud_file_context(db: Session, user_id: int, node: dict) -> str:
    file_ids = [value for value in node.get("uploaded_file_ids", []) if isinstance(value, int)]
    if not file_ids:
        return ""
    if len(file_ids) > MAX_FILES:
        raise CloudFileError(f"Cloud nodes support at most {MAX_FILES} attached files.")
    rows = db.query(models.UploadedFile).filter(
        models.UploadedFile.user_id == user_id,
        models.UploadedFile.id.in_(file_ids),
    ).all()
    by_id = {row.id: row for row in rows}
    parts = []
    for file_id in file_ids:
        uploaded = by_id.get(file_id)
        if uploaded is None:
            raise CloudFileError(f"Attached file {file_id} is unavailable.")
        path = UPLOAD_DIR / uploaded.stored_name
        if not path.is_file():
            raise CloudFileError(f"Stored file {uploaded.original_name} is unavailable.")
        content = extract_text(path)
        if not content.strip():
            raise CloudFileError(f"No extractable text in {uploaded.original_name}.")
        parts.append(f"[Attached file text: {uploaded.original_name}]\n{content}")
    return "\n\n".join(parts)
